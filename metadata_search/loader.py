"""Read the first two sheets without correcting original business records."""

from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from pathlib import Path
import re

from openpyxl import load_workbook


class CatalogError(ValueError):
    """The input cannot be read as a two-sheet metadata catalog."""


@dataclass(frozen=True)
class FieldRecord:
    """One source field row; source coordinates distinguish duplicate names."""

    system_name: str
    table_comment: str
    field_comment: str
    sample_data: str
    database_name: str
    table_name: str
    field_name: str
    source_file: str
    source_sheet: str
    source_row: int

    def to_dict(self) -> dict:
        return asdict(self)


# Keep source-language headers so existing Excel templates remain readable.
ALIASES = {
    "database_name": ("数据库名", "数据库英文名", "库名"),
    "table_name": ("数据表名", "表英文名", "表名"),
    "table_comment": ("业务数据资源名称", "表注释", "表中文名", "数据表注释"),
    "field_name": ("数据字段名", "字段英文名", "字段名"),
    "field_comment": ("数据字段注释", "字段注释", "字段中文名"),
    "sample_data": ("数据样例", "样例数据", "字段样例"),
    "sequence": ("序号",),
}


def cell_text(value: object) -> str:
    """Convert Excel values to display text without stripping their content."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value)


def _header_map(row: tuple) -> dict[str, int]:
    labels = [cell_text(value).strip() for value in row]
    return {
        key: labels.index(alias)
        for key, aliases in ALIASES.items()
        for alias in aliases
        if alias in labels
    }


def _business_rows(sheet, required: set[str]):
    """Locate a header, then skip template rows rather than filtering business data."""
    rows = sheet.iter_rows(values_only=True)
    columns = None
    for row_number, row in enumerate(rows, start=1):
        found = _header_map(row)
        if required <= found.keys():
            columns = found
            break
        if row_number >= 50:
            break
    if columns is None:
        names = ", ".join(ALIASES[key][0] for key in sorted(required))
        raise CatalogError(f"{sheet.title}: Required headers not found in the first 50 rows: {names}")

    def get(row, key):
        index = columns.get(key)
        return cell_text(row[index]) if index is not None and index < len(row) else ""

    for source_row, row in enumerate(rows, start=row_number + 1):
        if not any(cell_text(value).strip() for value in row):
            continue
        marker = get(row, "sequence").strip() if "sequence" in columns else cell_text(row[0]).strip()
        # These literal markers identify instructions and examples in the source template.
        if marker == "填表说明" or re.match(r"^(?:例|示例)\s*[:：]", marker):
            continue
        if required <= _header_map(row).keys():
            continue
        yield source_row, {key: get(row, key) for key in ALIASES}, "database_name" in columns


def load_workbook_records(path: Path) -> list[FieldRecord]:
    # data_only reads cached formula values; the workbook is never saved or changed.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) < 2:
            raise CatalogError(f"{path.name}: At least two worksheets are required")
        table_sheet, field_sheet = workbook.worksheets[:2]
        # Include the database in the join key to avoid mixing same-named tables.
        tables: dict[tuple[str, str], list[str]] = {}
        table_has_database = False
        for _, values, has_database in _business_rows(table_sheet, {"table_name", "table_comment"}):
            table_has_database = has_database
            key = (values["database_name"].strip(), values["table_name"].strip())
            comments = tables.setdefault(key, [])
            if values["table_comment"] not in comments:
                comments.append(values["table_comment"])

        records = []
        for row_number, values, _ in _business_rows(field_sheet, {"table_name", "field_name"}):
            database = values["database_name"].strip() if table_has_database else ""
            key = (database, values["table_name"].strip())
            # Keep every field row, even if its table metadata is missing.
            records.append(FieldRecord(
                system_name=path.stem,
                table_comment=";".join(tables.get(key, [])),
                field_comment=values["field_comment"],
                sample_data=values["sample_data"],
                database_name=values["database_name"],
                table_name=values["table_name"],
                field_name=values["field_name"],
                source_file=path.name,
                source_sheet=field_sheet.title,
                source_row=row_number,
            ))
        return records
    finally:
        workbook.close()


def load_catalog(directory: Path) -> tuple[list[FieldRecord], list[str]]:
    """Load files in a fixed order so equal-ranked results remain stable."""
    if not directory.is_dir():
        raise CatalogError(f"Data directory does not exist: {directory}")
    paths = sorted(
        (path for path in directory.iterdir()
         if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}
         and not path.name.startswith("~$")),
        key=lambda path: path.name,
    )
    if not paths:
        raise CatalogError(f"No .xlsx or .xlsm files found in: {directory}")
    records = []
    for path in paths:
        records.extend(load_workbook_records(path))
    return records, [path.name for path in paths]
